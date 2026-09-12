from io import BytesIO

from django.urls import reverse
from model_bakery import baker
from openpyxl import load_workbook

from bpp.models import Funkcja_Autora


def _ustaw_funkcje_autora(praca_z_dyscyplina):
    autor = praca_z_dyscyplina.autorzy_set.get().autor
    funkcja = baker.make(Funkcja_Autora, nazwa="kierownik zakładu", skrot="kz")
    autor.aktualna_funkcja = funkcja
    autor.save(update_fields=["aktualna_funkcja"])
    return funkcja.nazwa, praca_z_dyscyplina.typ_kbn.nazwa


def test_raport_slotow_ewaluacja_html_typ_kbn_i_funkcja_fd473(
    admin_client, praca_z_dyscyplina, rok
):
    funkcja, typ_kbn = _ustaw_funkcje_autora(praca_z_dyscyplina)

    res = admin_client.get(
        reverse("raport_slotow:raport-ewaluacja"),
        {"_export": "html", "od_roku": rok, "do_roku": rok},
    )

    assert res.status_code == 200
    table = res.context["table"]
    column_names = table.columns.names()
    assert column_names.index("typ_kbn") == column_names.index("rodzaj_publikacji") + 1
    assert column_names.index("funkcja") == column_names.index("autor") + 1
    row = next(iter(table.rows))
    assert row.get_cell_value("typ_kbn") == typ_kbn
    assert row.get_cell_value("funkcja") == funkcja


def test_raport_slotow_ewaluacja_xlsx_typ_kbn_i_funkcja_fd473(
    admin_client, praca_z_dyscyplina, rok
):
    funkcja, typ_kbn = _ustaw_funkcje_autora(praca_z_dyscyplina)

    res = admin_client.get(
        reverse("raport_slotow:raport-ewaluacja"),
        {"_export": "xlsx", "od_roku": rok, "do_roku": rok},
    )

    assert res.status_code == 200
    rows = list(load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True))
    headers = next(row for row in rows if "Autor ewaluowany" in row)
    assert "Typ MNiSW/MEiN" in headers
    assert "Funkcja" in headers
    values = {value for row in rows for value in row if value is not None}
    assert typ_kbn in values
    assert funkcja in values

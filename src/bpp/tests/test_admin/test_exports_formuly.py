from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse
from django_webtest import DjangoTestApp
from model_bakery import baker

from bpp.models import Autor

FORMULA = '=HYPERLINK("http://x","y")'


@pytest.mark.django_db
def test_xlsx_export_admin_zapisuje_tekst_zamiast_formuly(admin_app: DjangoTestApp):
    baker.make(Autor, nazwisko=FORMULA, imiona="Jan")

    page = admin_app.get(reverse("admin:bpp_autor_changelist"))
    res = page.click("Eksport")

    ws = openpyxl.load_workbook(BytesIO(res.content)).active
    naglowki = [c.value for c in ws[1]]
    komorka = ws.cell(row=2, column=naglowki.index("nazwisko") + 1)
    assert komorka.data_type == "s"
    assert komorka.value == "'" + FORMULA

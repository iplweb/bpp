import pytest
from django.contrib.auth.models import Group
from django.urls import reverse

from bpp.const import GR_WPROWADZANIE_DANYCH


def url(export_format, **params):
    return (
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": export_format})
        + "?"
        + "&".join(f"{k}={v}" for k, v in params.items())
    )


@pytest.fixture
def redaktor(client, admin_user):
    client.force_login(admin_user)
    return client


@pytest.mark.django_db
def test_eksport_csv_ma_naglowek_i_wiersz(redaktor, wydawnictwo_ciagle, denorms):
    # denorms.flush() to konwencja tego repo dla testów odpytujących Rekord
    # (patrz test_zapytanie.py) — pola "rok"/"tytul_oryginalny" są tu
    # faktycznie synchronizowane triggerem SQL od razu przy .save(), więc
    # flush nie jest ściśle wymagany, ale zostaje defensywnie na wypadek
    # przyszłych zmian testu na pola liczone przez django-denorm.
    denorms.flush()

    res = redaktor.get(
        url("csv", model="rekord", query=f"rok+%3D+{wydawnictwo_ciagle.rok}")
    )

    assert res.status_code == 200
    assert res["Content-Type"].startswith("text/csv")
    assert b"tytul_oryginalny" in res.content
    assert wydawnictwo_ciagle.tytul_oryginalny.encode() in res.content


@pytest.mark.django_db
def test_eksport_xlsx(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()

    res = redaktor.get(
        url("xlsx", model="rekord", query=f"rok+%3D+{wydawnictwo_ciagle.rok}")
    )

    assert "spreadsheetml" in res["Content-Type"]


@pytest.mark.django_db
def test_eksport_html_z_postaci_lista(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()

    res = redaktor.get(
        url(
            "html",
            model="rekord",
            query=f"rok+%3D+{wydawnictwo_ciagle.rok}",
            postac="list",
        )
    )

    assert res["Content-Type"].startswith("text/html")
    # sanitize_export_html (nh3) zdejmuje atrybut class — "multiseek-list-
    # report" istnieje tylko w renderze na stronie, nie w eksporcie.
    # Asercja na strukturze (<ol>, nie <table>) i realnej treści rekordu.
    assert b"<ol" in res.content
    assert b"<table" not in res.content
    assert wydawnictwo_ciagle.tytul_oryginalny.encode() in res.content


@pytest.mark.django_db
def test_eksport_bib_tylko_przy_postaci_bibtex(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()

    zle = redaktor.get(
        url(
            "bib",
            model="rekord",
            query=f"rok+%3D+{wydawnictwo_ciagle.rok}",
            postac="list",
        )
    )
    assert zle.status_code == 400

    dobrze = redaktor.get(
        url(
            "bib",
            model="rekord",
            query=f"rok+%3D+{wydawnictwo_ciagle.rok}",
            postac="bibtex",
        )
    )
    assert dobrze.status_code == 200
    assert "bibtex" in dobrze["Content-Type"]
    assert b"@" in dobrze.content


@pytest.mark.django_db
def test_eksport_nieznany_format_400(redaktor):
    res = redaktor.get(url("pdf", model="rekord", query="rok+%3D+2024"))
    assert res.status_code == 400


@pytest.mark.django_db
def test_eksport_bledne_zapytanie_400(redaktor):
    res = redaktor.get(url("csv", model="rekord", query="rok+%3D%3D%3D"))
    assert res.status_code == 400


@pytest.mark.django_db
def test_eksport_bledne_zapytanie_nie_odbija_niewyekranowanego_html(redaktor):
    """Reflected XSS: błąd DjangoQL odbija SUROWE literały zapytania.

    ``rok = "<script>...</script>"`` jest błędem typu (rok jest int, nie
    string) — komunikat wyjątku djangoql zawiera dosłowny literał z
    zapytania. Bez text/plain + escape() ten string wyrenderowałby się w
    przeglądarce jako aktywny <script> (jeden spreparowany link do
    zalogowanego redaktora/superusera — dokładnie tych, którzy mają tu
    dostęp).
    """
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "rekord",
            "query": 'rok = "<script>alert(1)</script>"',
        },
    )

    assert res.status_code == 400
    assert res["Content-Type"].startswith("text/plain")
    assert b"<script>alert(1)</script>" not in res.content


@pytest.mark.django_db
def test_eksport_anonim_403(client, wydawnictwo_ciagle):
    # WprowadzanieDanychOrSuperuserMixin ma raise_exception = True.
    # AccessMixin.handle_no_permission (Django) sprawdza `raise_exception OR
    # user.is_authenticated` — dla raise_exception=True warunek jest
    # prawdziwy NIEZALEZNIE od zalogowania, wiec anonim dostaje 403
    # (PermissionDenied), a nie redirect do loginu.
    res = client.get(url("csv", model="rekord", query="rok+%3D+2024"))
    assert res.status_code == 403


@pytest.mark.django_db
def test_eksport_staff_poza_grupa_403(client, django_user_model):
    user = django_user_model.objects.create_user(
        username="staff", password="x", is_staff=True
    )
    client.force_login(user)
    res = client.get(url("csv", model="rekord", query="rok+%3D+2024"))
    assert res.status_code == 403


@pytest.mark.django_db
def test_eksport_staff_w_grupie_ma_dostep(
    client, django_user_model, wydawnictwo_ciagle, denorms
):
    denorms.flush()
    user = django_user_model.objects.create_user(
        username="redaktor2", password="x", is_staff=True
    )
    group, _ = Group.objects.get_or_create(name=GR_WPROWADZANIE_DANYCH)
    user.groups.add(group)
    client.force_login(user)

    res = client.get(
        url("csv", model="rekord", query=f"rok+%3D+{wydawnictwo_ciagle.rok}")
    )
    assert res.status_code == 200


@pytest.mark.django_db
def test_eksport_sanityzuje_tytul_w_naglowku(redaktor, wydawnictwo_ciagle, denorms):
    """Tytuł jest w pełni user-controlled i trafia do Content-Disposition."""
    denorms.flush()

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "tytul": 'zły"tytuł\n../etc/passwd',
        },
    )

    disposition = res["Content-Disposition"]
    assert "\n" not in disposition
    assert "../" not in disposition


@pytest.mark.django_db
def test_eksport_dokumentu_powyzej_capu_400(
    redaktor, wydawnictwo_ciagle, denorms, monkeypatch
):
    from bpp.views import zapytanie_export

    # Denorms.flush() jest tu defensywny (patrz task-4-report.md — pole
    # "rok" jest synchronizowane triggerem SQL, nie wymaga flush), ale
    # zostaje dla zgodności z konwencją repo i odporności na przyszłe
    # zmiany zapytania na pole zależne od Pythonowego denorm.
    denorms.flush()
    monkeypatch.setattr(zapytanie_export, "ZAPYTANIE_EXPORT_MAX_DOKUMENT", 0)

    res = redaktor.get(
        url(
            "html",
            model="rekord",
            query=f"rok+%3D+{wydawnictwo_ciagle.rok}",
            postac="list",
        )
    )
    assert res.status_code == 400
    # Komunikat czyta limit ze stałej modułu (monkeypatch = 0), więc nie może
    # być zaszytego „5000" w tekście.
    assert b"maksymalnie 0 rekord" in res.content


# --- Zadanie 11: eksport LISTY autorów (CSV/XLSX) z metrykami dorobku ------


@pytest.mark.django_db
def test_eksport_autorow_csv_ma_kolumny_dorobku(redaktor, autor_jan_nowak):
    """Autor bez żadnej publikacji: brak `denorms.flush()` jest CELOWY (nie
    tworzymy tu żadnej publikacji, więc nie ma czego flushować) — wiersz ma
    się pojawić z liczba_prac=0, nie zniknąć i nie wysypać eksportu."""
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {"model": "autor", "query": 'nazwisko = "Nowak"'},
    )

    assert res.status_code == 200
    assert b"liczba_prac" in res.content
    assert b"suma_slotow" in res.content
    assert b"suma_pkdaut" in res.content
    assert b"Nowak" in res.content

    import csv
    import io

    wiersze = list(csv.reader(io.StringIO(res.content.decode())))
    naglowek, dane = wiersze[0], wiersze[1]
    kol = dict(zip(naglowek, dane, strict=True))
    assert kol["liczba_prac"] == "0"
    assert kol["suma_slotow"] == ""
    assert kol["suma_pkdaut"] == ""


@pytest.mark.django_db
def test_eksport_autorow_nie_zawyza_metryk(
    redaktor,
    zwarte_z_dyscyplinami,
    wydawnictwo_ciagle,
    autor_jan_nowak,
    jednostka,
    dyscyplina1,
    denorms,
):
    """Dwa JOIN-y do relacji do-wielu w jednym annotate zawyżyłyby sumy.

    Dane budowane PRODUKCYJNĄ ścieżką — dokładnie tak, jak
    test_baza_udzialow_nie_zawyza_sumy_przy_wielu_pracach w
    test_pivot_autor.py (ten sam scenariusz, żeby eksport listy i tabela
    krzyżowa autorów nie mogły się kiedyś rozjechać). Brief tego zadania miał
    tu fabrykowane `Cache_Punktacja_Autora(rekord_id=[1, 1])` wskazujące na
    nieistniejące rekordy — `cache_punktacja_autora_query` ma REALNY FK do
    Rekordu, więc taki JOIN albo się nie domyka, albo trafia gdzie popadnie
    i niczego by nie dowodził. Tutaj Jan Nowak ma DWIE prace i DWA wiersze
    udziału z realnymi rekord_id, wyliczone przez przelicz_punkty_dyscyplin():

    * `zwarte_z_dyscyplinami` — dwóch autorów, slot 0.5, pkdaut 10;
    * `wydawnictwo_ciagle` podniesione do 20 pkt, jeden autor (Nowak),
      slot 1.0, pkdaut 20.

    Poprawna Σ slotów = 0.5 + 1.0 = 1.5, Σ pkdaut = 10 + 20 = 30, liczba prac
    = 2. Zmierzone (patrz docstring test_baza_udzialow_nie_zawyza_sumy_przy_
    wielu_pracach): jeden annotate() na dwóch RÓŻNYCH relacjach do-wielu
    (`autorzy` i `cache_punktacja_autora_query`) daje iloczyn kartezjański
    2×2=4 i Σ slotów 3.0000 / Σ pkdaut 60.0000 — dokładnie dwukrotność.
    liczba_prac zostałaby poprawna (2), bo Count(distinct=True) jest na to
    odporny — tylko Sum zawyża, dlatego test sprawdza WSZYSTKIE TRZY liczby.
    """
    from decimal import Decimal

    from bpp.models.cache import Cache_Punktacja_Autora_Query

    wydawnictwo_ciagle.punkty_kbn = 20
    wydawnictwo_ciagle.save()
    wydawnictwo_ciagle.dodaj_autora(
        autor_jan_nowak, jednostka, dyscyplina_naukowa=dyscyplina1
    )
    wydawnictwo_ciagle.przelicz_punkty_dyscyplin()
    denorms.flush()

    # Dowód, że JOIN przez cache_punktacja_autora_query naprawdę trafia w
    # istniejące Rekordy (gdyby rekord_id nie pasował, sumy byłyby zerowe, a
    # test „zielony z pustki") — dwa wiersze udziału, dwie prace w autorzy.
    assert (
        Cache_Punktacja_Autora_Query.objects.filter(autor=autor_jan_nowak).count() == 2
    )
    assert autor_jan_nowak.autorzy_set.count() == 2

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {"model": "autor", "query": 'nazwisko = "Nowak"'},
    )

    import csv
    import io

    wiersze = list(csv.reader(io.StringIO(res.content.decode())))
    naglowek, dane = wiersze[0], wiersze[1]
    kol = dict(zip(naglowek, dane, strict=True))

    # Bez rozdzielenia agregatów na dwa zapytania Σ slotów wyszłoby 3.0000, a
    # Σ pkdaut 60.0000 (patrz docstring wyżej) — dokładne liczby, nie „> 0".
    assert int(kol["liczba_prac"]) == 2
    assert Decimal(kol["suma_slotow"]) == Decimal("1.5000")
    assert Decimal(kol["suma_pkdaut"]) == Decimal("30.0000")


@pytest.mark.django_db
def test_eksport_autorow_xlsx(
    redaktor, zwarte_z_dyscyplinami, autor_jan_nowak, denorms
):
    """Kontrapunkt content-type-only smoke testu: otwiera realny workbook
    (openpyxl) i sprawdza nagłówek + komórkę danych — pusty/zepsuty plik XLSX
    ma też poprawny Content-Type, więc samo to niczego nie dowodzi."""
    import io

    from openpyxl import load_workbook

    from bpp.views.multiseek_export import AUTOR_EXPORT_XLSX_HEADERS

    denorms.flush()

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "xlsx"}),
        {"model": "autor", "query": 'nazwisko = "Nowak"'},
    )

    assert res.status_code == 200
    assert "spreadsheetml" in res["Content-Type"]

    workbook = load_workbook(io.BytesIO(res.content))
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0] == AUTOR_EXPORT_XLSX_HEADERS
    assert len(rows) == 2

    kol = dict(zip(rows[0], rows[1], strict=True))
    assert kol["Nazwisko"] == "Nowak"
    assert kol["Liczba prac"] == 1
    assert kol["Σ slotów"] == 0.5
    assert kol["Σ pkdaut"] == 10.0

    # Σ slotów/Σ pkdaut MUSZĄ być liczbami (data_type "n"), nie tekstem —
    # inaczej SUM() nad kolumną w Excelu policzy 0, a nie realną sumę.
    slot_col = AUTOR_EXPORT_XLSX_HEADERS.index("Σ slotów") + 1
    pkdaut_col = AUTOR_EXPORT_XLSX_HEADERS.index("Σ pkdaut") + 1
    slot_cell = worksheet.cell(row=2, column=slot_col)
    pkdaut_cell = worksheet.cell(row=2, column=pkdaut_col)
    assert slot_cell.data_type == "n"
    assert pkdaut_cell.data_type == "n"
    assert slot_cell.number_format == "0.0000"
    assert pkdaut_cell.number_format == "0.0000"


@pytest.mark.django_db
def test_eksport_autorow_html_400(redaktor, autor_jan_nowak):
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "html"}),
        {"model": "autor", "query": 'nazwisko = "Nowak"'},
    )

    assert res.status_code == 400

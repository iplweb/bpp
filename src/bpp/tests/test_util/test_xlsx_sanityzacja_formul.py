"""Sanityzacja komórek tekstowych w eksportach arkuszy (XLSX/CSV)."""

import datetime
from decimal import Decimal
from io import BytesIO

import django_tables2 as tables
import openpyxl
import pytest
import tablib

from bpp.util.xlsx import (
    sanitize_tablib_dataset,
    sanitize_xlsx_cell,
    sanitize_xlsx_row,
    sanitize_xlsx_workbook,
    sanitize_xlsx_worksheet,
)

FORMULA = '=HYPERLINK("http://x","y")'


@pytest.mark.parametrize(
    "value",
    [
        FORMULA,
        "=1+1",
        "+1+1",
        "-1+1",
        "@SUM(A1)",
        "\tfoo",
        "\rfoo",
        "\nfoo",
    ],
)
def test_sanitize_xlsx_cell_prefiksuje_apostrofem(value):
    assert sanitize_xlsx_cell(value) == "'" + value


@pytest.mark.parametrize(
    "value",
    [
        None,
        0,
        -5,
        12.5,
        Decimal("-1.25"),
        datetime.date(2020, 1, 1),
        "",
        "Tytuł",
        "-",
        "-12",
        "-12,5",
        "+3.25",
        "a=b",
    ],
)
def test_sanitize_xlsx_cell_nie_zmienia_bezpiecznych(value):
    assert sanitize_xlsx_cell(value) == value


def test_sanitize_xlsx_cell_idempotentne():
    once = sanitize_xlsx_cell(FORMULA)
    assert sanitize_xlsx_cell(once) == once


def test_sanitize_xlsx_row():
    assert sanitize_xlsx_row([FORMULA, 1, None]) == ["'" + FORMULA, 1, None]


def test_sanitize_tablib_dataset():
    ds = tablib.Dataset(headers=["Tytuł", "Punkty"])
    ds.append([FORMULA, 10])
    ds.append(["ok", Decimal("2.5")])
    sanitize_tablib_dataset(ds)
    assert ds.headers == ["Tytuł", "Punkty"]
    assert list(ds[0]) == ["'" + FORMULA, 10]
    assert list(ds[1]) == ["ok", Decimal("2.5")]


def _zapisz_i_wczytaj(wb):
    stream = BytesIO()
    wb.save(stream)
    return openpyxl.load_workbook(BytesIO(stream.getvalue()))


def test_sanitize_xlsx_worksheet_zamienia_formule_na_tekst():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([FORMULA, 5, "-"])
    assert ws["A1"].data_type == "f"
    sanitize_xlsx_worksheet(ws)

    ws = _zapisz_i_wczytaj(wb).active
    assert ws["A1"].data_type == "s"
    assert ws["A1"].value == "'" + FORMULA
    assert ws["B1"].value == 5
    assert ws["C1"].value == "-"


def test_sanitize_xlsx_workbook_wszystkie_arkusze():
    wb = openpyxl.Workbook()
    wb.active.append([FORMULA])
    wb.create_sheet("Drugi").append([FORMULA])
    sanitize_xlsx_workbook(wb)

    wb = _zapisz_i_wczytaj(wb)
    for ws in wb.worksheets:
        assert ws["A1"].data_type == "s"
        assert ws["A1"].value == "'" + FORMULA


def test_prettyxlsx_zapisuje_tekst_a_url_jako_hyperlink():
    from bpp.admin.xlsx_export.formats import PrettyXLSX

    ds = tablib.Dataset(headers=["tytul", "punkty", "bpp_strona_url"])
    ds.append([FORMULA, 10, "https://bpp.example/rekord/1/"])
    ws = openpyxl.load_workbook(BytesIO(PrettyXLSX().export_data(ds))).active

    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'" + FORMULA
    assert ws["B2"].value == 10
    assert ws["C2"].data_type == "f"
    assert ws["C2"].value.startswith('=HYPERLINK("https://bpp.example/')


class _Tabela(tables.Table):
    tytul = tables.Column()
    punkty = tables.Column()


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_bezpieczny_table_export_arkusze(fmt):
    from bpp.export.tables2 import BezpiecznyTableExport

    exp = BezpiecznyTableExport(fmt, _Tabela([{"tytul": FORMULA, "punkty": 7}]))
    assert list(exp.dataset[0]) == ["'" + FORMULA, 7]


def test_bezpieczny_table_export_xlsx_komorka_tekstowa():
    from bpp.export.tables2 import BezpiecznyTableExport

    data = BezpiecznyTableExport(
        "xlsx", _Tabela([{"tytul": FORMULA, "punkty": 7}])
    ).export()
    ws = openpyxl.load_workbook(BytesIO(data)).active
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'" + FORMULA
    assert ws["B2"].value == 7


def test_bezpieczny_table_export_json_bez_zmian():
    from bpp.export.tables2 import BezpiecznyTableExport

    exp = BezpiecznyTableExport("json", _Tabela([{"tytul": FORMULA, "punkty": 7}]))
    assert list(exp.dataset[0]) == [FORMULA, 7]


def test_raport_slotow_my_table_export_dane_i_opis_jako_tekst():
    from raport_slotow.util import MyTableExport

    class _Demo(tables.Table):
        name = tables.Column()
        value = tables.Column(footer=lambda table: 0)

    data = MyTableExport(
        "xlsx",
        _Demo([{"name": FORMULA, "value": 10}]),
        export_description=[("Autor:", FORMULA), [FORMULA]],
    ).export_xlsx()
    ws = openpyxl.load_workbook(BytesIO(data)).active
    rows = list(ws.iter_rows())

    # Opis
    assert rows[0][1].data_type == "s"
    assert rows[0][1].value == "'" + FORMULA
    assert rows[1][0].value == "'" + FORMULA
    # Dane (po pustym wierszu i nagłówku)
    dane = rows[4]
    assert dane[0].data_type == "s"
    assert dane[0].value == "'" + FORMULA
    assert dane[1].value == 10
    # Stopka nadal jest formułą
    assert rows[5][1].data_type == "f"
    assert rows[5][1].value.startswith("=SUBTOTAL(")

"""Tabela krzyżowa (pivot) na stronie "Szukaj zapytaniem" (postac=pivot).

Silnik pivota (parse_pivot_params/zbuduj_pivot/DIMENSIONS/METRICS) jest
gotowy i przetestowany w test_multiseek_pivot.py — tu sprawdzamy TYLKO
podpięcie: kontekst widoku, render partiala i eksport macierzy CSV/XLSX
z poziomu /zapytanie/.
"""

import csv
import io
import re

import pytest
from django.urls import reverse


@pytest.fixture
def redaktor(client, admin_user):
    client.force_login(admin_user)
    return client


def _wiersze_csv(response):
    """CSV eksportu jako lista list — asercje na TREŚĆ macierzy, nie na samo
    słowo „RAZEM" (które jest w KAŻDEJ macierzy, także w tej domyślnej, więc
    nie odróżnia poprawnego eksportu od zignorowania pivot_row/col/val)."""
    return list(csv.reader(io.StringIO(response.content.decode())))


def _linki_eksportu(response):
    """Wszystkie linki eksportu wyrenderowane na stronie (z rozkodowanym
    &amp;) — górny pasek i pasek pod macierzą razem."""
    return [
        href.replace("&amp;", "&")
        for href in re.findall(
            r'href="([^"]*/zapytanie/eksport/[^"]*)"', response.content.decode()
        )
    ]


@pytest.mark.django_db
def test_pivot_rekordow_renderuje_macierz(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_val": "liczba",
        },
    )
    assert res.status_code == 200
    assert res.context["pivot"].grand_total == 1
    assert str(wydawnictwo_ciagle.rok).encode() in res.content


@pytest.mark.django_db
def test_pivot_rekordow_z_kolumna(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_col": "typ_kbn",
            "pivot_val": "liczba",
        },
    )
    assert res.context["pivot"].col_dim.key == "typ_kbn"


@pytest.mark.django_db
def test_pivot_formularz_przenosi_stan_strony(redaktor, wydawnictwo_ciagle, denorms):
    """P1/P3: bez ukrytych pól model/query/postac submit zgubiłby filtr —
    strona zapytania trzyma go w URL-u, nie w sesji (jak multiseek)."""
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_val": "liczba",
        },
    )
    content = res.content.decode()
    assert f'action="{reverse("bpp:zapytanie")}"' in content
    assert 'name="model" value="rekord"' in content
    assert f'name="query" value="rok = {wydawnictwo_ciagle.rok}"' in content
    assert 'name="postac" value="pivot"' in content


@pytest.mark.django_db
def test_pivot_linki_eksportu_pod_wlasciwym_urlem(
    redaktor, wydawnictwo_ciagle, denorms
):
    """P2/P3: linki eksportu muszą wskazywać /zapytanie/eksport/, nie
    ../export/ (ta ścieżka istnieje tylko pod /multiseek/results/)."""
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_val": "liczba",
        },
    )
    content = res.content.decode()
    base = reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"})
    base = base.rsplit("csv/", 1)[0]
    assert f"{base}xlsx/?" in content
    assert f"{base}csv/?" in content
    assert "../export/" not in content


@pytest.mark.django_db
def test_pivot_zbyt_duza_macierz_pokazuje_komunikat(
    redaktor, wydawnictwo_ciagle, denorms, monkeypatch
):
    """Bramka rozmiaru (PIVOT_MAX_CELLS) jest podłączona pod /zapytanie/ —
    zamiast wywalonego 500 user dostaje czytelny komunikat na ekranie."""
    from bpp.pivot import core as pivot_core

    # Prog czyta zbuduj_pivot jako globalna SWOJEGO modulu, wiec patchujemy go
    # w bpp.pivot.core, a nie przez zgodnosciowy shim multiseek_registry.pivot.
    monkeypatch.setattr(pivot_core, "PIVOT_MAX_CELLS", 0)
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_val": "liczba",
        },
    )
    assert res.status_code == 200
    assert res.context["pivot"] is None
    assert b"zbyt du" in res.content


@pytest.mark.django_db
def test_pivot_eksport_csv_zwraca_ZAMOWIONA_macierz(
    redaktor, wydawnictwo_ciagle, denorms
):
    """Eksport MUSI zbudować macierz z parametrów GET, nie domyślną.

    Asercja na samo `b"RAZEM"` niczego nie dowodziła: to słowo jest w KAŻDEJ
    macierzy, także w domyślnej („Rok,RAZEM / 2026,1,1 / RAZEM,1"). Dlatego
    tutaj sprawdzamy nagłówek zależny od WYBORU (etykieta wymiaru wiersza i
    kolumny) oraz dokładną wartość komórki, plus kontrapunkt: nagłówka
    macierzy domyślnej ma NIE być.
    """
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "charakter_ogolny",
            "pivot_col": "typ_kbn",
            "pivot_val": "liczba",
        },
    )
    assert res.status_code == 200

    wiersze = _wiersze_csv(res)
    assert wiersze[0] == [
        "Charakter ogólny (rodzaj)",
        str(wydawnictwo_ciagle.typ_kbn),
        "RAZEM",
    ]
    assert wiersze[-1] == ["RAZEM", "1", "1"]
    assert b"Rok,RAZEM" not in res.content


@pytest.mark.django_db
def test_pivot_eksport_xlsx_zwraca_ZAMOWIONA_macierz(
    redaktor, wydawnictwo_ciagle, denorms
):
    """To samo co wyżej, ale przez openpyxl — sam Content-Type nie mówi nic
    o tym, KTÓRA macierz wylądowała w skoroszycie."""
    from openpyxl import load_workbook

    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "xlsx"}),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "charakter_ogolny",
            "pivot_col": "typ_kbn",
            "pivot_val": "liczba",
        },
    )
    assert res.status_code == 200
    assert "spreadsheetml" in res["Content-Type"]

    arkusz = load_workbook(io.BytesIO(res.content)).active
    wiersze = [[c.value for c in row] for row in arkusz.iter_rows()]
    assert wiersze[0] == [
        "Charakter ogólny (rodzaj)",
        str(wydawnictwo_ciagle.typ_kbn),
        "RAZEM",
    ]
    assert wiersze[-1] == ["RAZEM", 1, 1]


@pytest.mark.django_db
def test_kazdy_link_eksportu_na_stronie_pivota_daje_WIDOCZNA_macierz(
    redaktor, wydawnictwo_ciagle, denorms
):
    """Regresja K1: strona pokazywała DWA przyciski CSV dające różne pliki.

    Górny pasek eksportu składał URL z samych `model`/`query`/`postac`, więc
    przy `postac=pivot` gubił `pivot_row`/`pivot_col`/`pivot_val` i ściągał
    macierz DOMYŚLNĄ (zmierzone: „Rok,RAZEM / 2026,1,1 / RAZEM,1"), podczas
    gdy pasek pod macierzą — budowany z `request.GET.urlencode` — ściągał tę
    właściwą. Test nie zakłada, KTÓRY pasek zostanie: bierze wszystkie linki
    eksportu obecne na stronie i wymaga, żeby KAŻDY z nich niósł wybór usera
    i zwracał macierz zgodną z ekranem.
    """
    denorms.flush()
    params = {
        "model": "rekord",
        "query": f"rok = {wydawnictwo_ciagle.rok}",
        "postac": "pivot",
        "pivot_row": "charakter_ogolny",
        "pivot_col": "typ_kbn",
        "pivot_val": "liczba",
    }
    strona = redaktor.get(reverse("bpp:zapytanie"), params)
    assert strona.status_code == 200
    oczekiwany_naglowek = strona.context["pivot"].as_table()["row_header"]
    assert oczekiwany_naglowek == "Charakter ogólny (rodzaj)"

    linki = _linki_eksportu(strona)
    assert linki, "macierz bez ŻADNEGO wejścia do eksportu to też regresja"

    for link in linki:
        assert "pivot_row=charakter_ogolny" in link, link
        assert "pivot_col=typ_kbn" in link, link
        if "/csv/" not in link:
            continue
        pobrane = redaktor.get(link)
        assert pobrane.status_code == 200, link
        assert _wiersze_csv(pobrane)[0][0] == oczekiwany_naglowek, link


@pytest.mark.django_db
def test_pivot_eksport_html_400(redaktor, wydawnictwo_ciagle, denorms):
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "html"}),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
        },
    )
    assert res.status_code == 400
    # 400 pivota idzie przez _blad() jak każdy inny błąd tego widoku —
    # text/plain chroni przed reflected XSS (patrz komentarz w
    # zapytanie_export.py przy _blad).
    assert res["Content-Type"].startswith("text/plain")


@pytest.mark.django_db
def test_pivot_w_liscie_postaci_rekordowych(redaktor):
    """Opcja "tabela krzyżowa" wraca do <select> — Zadanie 3 ją usunęło,
    bo bez implementacji renderowała placeholder listy."""
    from bpp.views.zapytanie import POSTAC_PIVOT, POSTACIE_REKORD

    assert POSTAC_PIVOT in {key for key, _ in POSTACIE_REKORD}


# ---------------------------------------------------------------------------
# Zadanie 9 — pivot autorski w UI + presety.
#
# Silnik (zbuduj_pivot_autora/DIMENSIONS/METRICS) jest gotowy i przetestowany
# w test_pivot_autor.py — tu sprawdzamy TYLKO podpięcie pod /zapytanie/:
# rozgałęzienie _pivot_context po modelu (wybierz_rejestr_pivota), eksport
# macierzy autorskiej i presety w kontekście strony.
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_pivot_autorow_jednostka_x_tytul(redaktor, jednostka):
    """Brak fixture'a `tytul` (liczba pojedyncza) w tym repo — jest tylko
    `tytuly` (mnoga). Tytuł budujemy więc inline przez baker.make, tak jak
    już robi to test_pivot_autor.py::test_krzyzowo_jednostka_x_tytul."""
    from model_bakery import baker

    from bpp.models import Autor
    from bpp.models.autor import Tytul

    tytul = baker.make(Tytul)
    baker.make(
        Autor,
        nazwisko="Nowak",
        aktualna_jednostka=jednostka,
        tytul=tytul,
        _quantity=2,
    )

    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_row": "jednostka",
            "pivot_col": "tytul",
            "pivot_val": "liczba_autorow",
        },
    )

    assert res.status_code == 200
    assert res.context["pivot"].grand_total == 2


@pytest.mark.django_db
def test_pivot_autorow_eksport_csv(redaktor, jednostka):
    from model_bakery import baker

    from bpp.models import Autor

    baker.make(Autor, nazwisko="Nowak", aktualna_jednostka=jednostka)

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_row": "jednostka",
            "pivot_val": "liczba_autorow",
        },
    )

    assert res.status_code == 200
    # Treść, nie samo "RAZEM": etykieta zamówionego wymiaru + dokładna komórka.
    # Bez wymiaru kolumn wiersze danych mają jedną kolumnę więcej niż nagłówek
    # (bezimienna kolumna wartości + RAZEM) — tak renderuje też ekran, patrz
    # PivotResult.as_table.
    assert _wiersze_csv(res) == [
        ["Aktualna jednostka", "RAZEM"],
        [str(jednostka), "1", "1"],
        ["RAZEM", "1"],
    ]


@pytest.mark.django_db
def test_pivot_autorow_eksport_xlsx(redaktor, jednostka):
    """Kontrapunkt CSV powyżej — XLSX to drugi format, który DEFEKT #4
    (eksport_formaty zwracające pustą krotkę dla autora) też chronił."""
    from model_bakery import baker

    from bpp.models import Autor

    baker.make(Autor, nazwisko="Nowak", aktualna_jednostka=jednostka)

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "xlsx"}),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_row": "jednostka",
            "pivot_val": "liczba_autorow",
        },
    )

    assert res.status_code == 200
    assert "spreadsheetml" in res["Content-Type"]


@pytest.mark.django_db
def test_pivot_autorow_lista_teraz_dziala(redaktor, autor_jan_nowak, denorms):
    """Kontrapunkt historyczny: do Zadania 11 eksport LISTY autorów (postac
    domyślna "rekordy") był zablokowany 400-ką — to rozróżnienie chroniło
    Zadanie 9 (macierz autorska) przed przypadkowym „otwarciem wszystkiego"
    rozgałęzieniem po postaci w ZapytanieExportView.get(). Od Zadania 11
    lista też działa naprawdę (autor_csv_export_response), więc 400 by tu był
    regresją — patrz test_zapytanie_export.py::test_eksport_autorow_csv_ma_
    kolumny_dorobku dla właściwego pokrycia tej ścieżki."""
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {"model": "autor", "query": 'nazwisko = "Nowak"'},
    )
    assert res.status_code == 200
    assert "text/csv" in res["Content-Type"]


@pytest.mark.django_db
def test_presety_bez_zapytania_sa_widoczne_ale_NIE_klikalne(redaktor):
    """Na świeżej stronie preset nie ma czego dokładać — i nie udaje, że ma.

    Do tej poprawki presety renderowały się jako linki z `query=` (pustym).
    Zmierzone kliknięcie takiego linku: 200, `results = None`, `pivot` w ogóle
    nieobecny w kontekście, zero `multiseek-pivot` w treści — strona
    przeładowywała się i NIC się nie działo, mimo że summary obiecuje „jeden
    klik do macierzy". Puste zapytanie nie przechodzi przez parser DjangoQL
    („Unexpected end of input"), więc nie da się go uratować po stronie
    serwera. Oferta zostaje widoczna (discoverability), ale bez linku.
    """
    res = redaktor.get(reverse("bpp:zapytanie"), {"model": "autor"})
    assert b"Audyt kompletno" in res.content
    assert b"po wykonaniu zapytania" in res.content
    assert not [
        href
        for href in re.findall(r'href="([^"]*)"', res.content.decode())
        if "postac=pivot" in href
    ]


@pytest.mark.django_db
def test_preset_po_wyslanym_zapytaniu_prowadzi_do_MACIERZY(redaktor, jednostka):
    """Kontrapunkt: gdy zapytanie jest wysłane, preset naprawdę robi to, co
    obiecuje — jeden klik i na ekranie stoi tabela krzyżowa."""
    from model_bakery import baker

    from bpp.models import Autor

    baker.make(Autor, nazwisko="Nowak", aktualna_jednostka=jednostka)

    strona = redaktor.get(
        reverse("bpp:zapytanie"), {"model": "autor", "query": 'nazwisko = "Nowak"'}
    )
    linki = [
        href.replace("&amp;", "&")
        for href in re.findall(r'href="([^"]*)"', strona.content.decode())
        if "postac=pivot" in href
    ]
    assert linki, "po wysłanym zapytaniu presety MUSZĄ być klikalne"

    wynik = redaktor.get(linki[0])
    assert wynik.status_code == 200
    assert wynik.context["pivot"] is not None
    assert wynik.context["pivot"].grand_total == 1
    assert b"multiseek-pivot" in wynik.content


@pytest.mark.django_db
def test_presety_dla_autora_niosa_biezace_zapytanie(redaktor):
    """Presety mają DOŁOŻYĆ wymiary/metrykę do query-stringa, nie zgubić
    zapytania, które user już wpisał — inaczej klik w preset cofnąłby go do
    pustego wyniku."""
    from django.template.defaultfilters import urlencode as tpl_urlencode

    from bpp.views.zapytanie import PIVOT_PRESETY_AUTOR

    biezace_zapytanie = 'nazwisko = "Nowak"'
    res = redaktor.get(
        reverse("bpp:zapytanie"), {"model": "autor", "query": biezace_zapytanie}
    )
    fragment = f"query={tpl_urlencode(biezace_zapytanie)}".encode()
    assert res.content.count(fragment) >= len(PIVOT_PRESETY_AUTOR)


@pytest.mark.django_db
def test_pivot_autorow_baza_udzialow_renderuje_sume(
    redaktor, zwarte_z_dyscyplinami, autor_jan_nowak, denorms
):
    """End-to-end bazy U na stronie: Σ slotów × rok renderuje się i eksportuje.

    `zwarte_z_dyscyplinami` buduje realne wiersze `bpp_cache_punktacja_autora`
    (dwóch autorów po 0.5 slota), więc macierz ma się z czego wziąć — grand
    total 1.0 dla obu autorów razem, 0.5 w wierszu Jana Nowaka.
    """
    from decimal import Decimal

    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko in ("Nowak", "Kowalski")',
            "postac": "pivot",
            "pivot_row": "autor",
            "pivot_col": "rok",
            "pivot_val": "suma_slotow",
        },
    )
    assert res.status_code == 200
    pivot = res.context["pivot"]
    assert pivot.grand_total == Decimal("1.0000")
    assert pivot.row_totals[autor_jan_nowak.pk] == Decimal("0.5000")
    assert str(zwarte_z_dyscyplinami.rok).encode() in res.content

    eksport = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "autor",
            "query": 'nazwisko in ("Nowak", "Kowalski")',
            "postac": "pivot",
            "pivot_row": "autor",
            "pivot_col": "rok",
            "pivot_val": "suma_slotow",
        },
    )
    assert eksport.status_code == 200
    wiersze = _wiersze_csv(eksport)
    # Nagłówek NIESIE wybór: wiersz „Autor" (nie domyślna „Aktualna
    # jednostka"), kolumna = rok publikacji, komórka = Σ slotów.
    assert wiersze[0] == ["Autor", str(zwarte_z_dyscyplinami.rok), "RAZEM"]
    assert wiersze[-1] == ["RAZEM", "1.0000", "1.0000"]
    assert [w for w in wiersze if w[0] == str(autor_jan_nowak)] == [
        [str(autor_jan_nowak), "0.5000", "0.5000"]
    ]


@pytest.mark.django_db
def test_presety_nieobecne_dla_rekordu(redaktor, wydawnictwo_ciagle, denorms):
    """Presety z Zadania 9 dotyczą TYLKO modelu autor (baza K, patrz brief).
    Model rekord ich nie pokazuje — zapobiega martwym linkom z parametrami
    wymiarów, które w rejestrze rekordowym nie istnieją (np. "ma_orcid")."""
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {"model": "rekord", "query": f"rok = {wydawnictwo_ciagle.rok}"},
    )
    assert b"Audyt kompletno" not in res.content


@pytest.mark.django_db
def test_metryka_kadrowa_pokazuje_rejestr_autorski(redaktor, autor_jan_nowak):
    """Metryka bazy K (liczba autorów) chowa wymiary publikacyjne baz P/U.

    Selektor dostaje z widoku PRZEFILTROWANY słownik wymiarów: sam `bpp_autor`
    nie ma dojścia do rekordu, więc „rok"/„dyscyplina"/… nie mają w bazie K
    ścieżki ORM (`expr_dla("K") is None`) i muszą zniknąć z selektora. Zostają
    dokładnie wymiary autorskie — porównanie z jawną listą, a nie z
    `DIMENSIONS.keys()`, bo od Zadania 10 rejestr jest szerszy niż baza K.
    """
    from bpp.pivot import autor as autor_rejestr

    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_val": "liczba_autorow",
        },
    )
    widoczne = res.context["pivot_dimensions"]
    assert set(widoczne) == {
        klucz
        for klucz, dim in autor_rejestr.DIMENSIONS.items()
        if dim.expr_dla(autor_rejestr.BAZA_KADROWA) is not None
    }
    assert "jednostka" in widoczne
    assert "rok" not in widoczne
    assert "dyscyplina" not in widoczne


@pytest.mark.django_db
def test_metryka_slotowa_pokazuje_rok(redaktor, autor_jan_nowak):
    """Metryka Σ slotów (baza U) odsłania wymiary publikacyjne bazy U — ale
    tylko te, które `bpp_cache_punktacja_autora` naprawdę zna: typ
    odpowiedzialności i charakter formalny zostają wyłącznie w bazie P."""
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_val": "suma_slotow",
        },
    )
    widoczne = res.context["pivot_dimensions"]
    assert "rok" in widoczne
    assert "dyscyplina" in widoczne
    assert "jednostka_pracy" in widoczne
    assert "typ_odpowiedzialnosci" not in widoczne
    assert "charakter_formalny" not in widoczne


@pytest.mark.django_db
def test_metryka_liczby_prac_pokazuje_wymiary_bazy_prac(redaktor, autor_jan_nowak):
    """Baza P (liczba prac) idzie przez `autorzy` do rekordu, więc widzi
    WSZYSTKIE wymiary publikacyjne — w odróżnieniu od bazy U."""
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_val": "liczba_prac",
        },
    )
    widoczne = res.context["pivot_dimensions"]
    assert {"rok", "dyscyplina", "typ_odpowiedzialnosci", "charakter_formalny"} <= set(
        widoczne
    )


@pytest.mark.django_db
def test_presety_autora_sa_wykonalne_w_swojej_bazie():
    """Każdy preset musi PRZEŻYĆ parse_pivot_params_autor bez podmiany.

    Preset podaje wiersz, kolumnę i metrykę; metryka wybiera bazę, a baza
    decyduje, które wymiary istnieją. Gdyby preset łączył wymiar publikacyjny
    (np. „rok") z metryką bazy kadrowej, parser cicho podstawiłby DEFAULT_ROW /
    wyrzucił kolumnę i user zobaczyłby tabelę inną niż obiecuje etykieta.
    """
    from bpp.pivot.autor import parse_pivot_params_autor
    from bpp.views.zapytanie import PIVOT_PRESETY_AUTOR

    for opis, row_key, col_key, metric_key in PIVOT_PRESETY_AUTOR:
        row, col, metric = parse_pivot_params_autor(
            {"pivot_row": row_key, "pivot_col": col_key, "pivot_val": metric_key}
        )
        assert metric.key == metric_key, opis
        assert row.key == row_key, opis
        assert col is not None and col.key == col_key, opis


# --- stronicowanie i sortowanie macierzy (/zapytanie/) ---------------------


def _autorzy_nowak(ile, jednostka):
    from model_bakery import baker

    from bpp.models import Autor

    return [
        baker.make(
            Autor,
            nazwisko="Nowak",
            imiona=f"Jan {i:03d}",
            aktualna_jednostka=jednostka,
        )
        for i in range(ile)
    ]


@pytest.mark.django_db
def test_pivot_autorow_stronicuje_wiersze(redaktor, jednostka):
    """Sedno zgłoszenia: pivot_row=autor przy 2 tys. autorów wypluwał
    wszystkie wiersze naraz."""
    _autorzy_nowak(30, jednostka)

    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_row": "autor",
            "pivot_val": "liczba_autorow",
            "pivot_per_page": "25",
        },
    )

    assert res.status_code == 200
    t = res.context["pivot_table"]
    assert t["wszystkich_wierszy"] == 30
    assert len(t["rows"]) == 25
    assert t["is_paginated"] is True
    assert "pivot_page=2" in res.content.decode().replace("&amp;", "&")


@pytest.mark.django_db
def test_pivot_druga_strona_pokazuje_reszte(redaktor, jednostka):
    _autorzy_nowak(30, jednostka)

    wspolne = {
        "model": "autor",
        "query": 'nazwisko = "Nowak"',
        "postac": "pivot",
        "pivot_row": "autor",
        "pivot_val": "liczba_autorow",
        "pivot_per_page": "25",
    }
    pierwsza = redaktor.get(reverse("bpp:zapytanie"), wspolne)
    druga = redaktor.get(reverse("bpp:zapytanie"), {**wspolne, "pivot_page": "2"})

    assert len(druga.context["pivot_table"]["rows"]) == 5
    etykiety_1 = {r["label"] for r in pierwsza.context["pivot_table"]["rows"]}
    etykiety_2 = {r["label"] for r in druga.context["pivot_table"]["rows"]}
    assert not (etykiety_1 & etykiety_2)
    # RAZEM nadal dotyczy całego zbioru, nie widocznej strony
    assert druga.context["pivot_table"]["grand_total"] == 30


@pytest.mark.django_db
def test_pivot_sortowanie_po_sumie_zmienia_kolejnosc(
    redaktor, wydawnictwo_ciagle, denorms
):
    """Wiersze mają iść wg kolumny RAZEM, nie alfabetycznie."""
    denorms.flush()
    wspolne = {
        "model": "rekord",
        "query": f"rok >= {wydawnictwo_ciagle.rok - 5}",
        "postac": "pivot",
        "pivot_row": "charakter_formalny",
        "pivot_val": "liczba",
    }
    res = redaktor.get(
        reverse("bpp:zapytanie"), {**wspolne, "pivot_sort": "suma", "pivot_dir": "desc"}
    )

    assert res.status_code == 200
    t = res.context["pivot_table"]
    sumy = [r["total"] or 0 for r in t["rows"]]
    assert sumy == sorted(sumy, reverse=True)
    assert t["sort"] == "suma"
    assert t["sort_suma_strzalka"] == "▼"


@pytest.mark.django_db
def test_pivot_formularz_przenosi_sortowanie(redaktor, wydawnictwo_ciagle, denorms):
    """Selecty auto-submitują; bez ukrytych pól zmiana wymiaru cicho
    wracałaby do kolejności domyślnej."""
    denorms.flush()
    res = redaktor.get(
        reverse("bpp:zapytanie"),
        {
            "model": "rekord",
            "query": f"rok = {wydawnictwo_ciagle.rok}",
            "postac": "pivot",
            "pivot_row": "rok",
            "pivot_val": "liczba",
            "pivot_sort": "suma",
            "pivot_dir": "asc",
        },
    )
    content = res.content.decode()
    assert 'name="pivot_sort" value="suma"' in content
    assert 'value="asc"' in content
    assert 'name="pivot_per_page"' in content


@pytest.mark.django_db
def test_pivot_eksport_ignoruje_strone_ale_bierze_sortowanie(redaktor, jednostka):
    """Plik ma zawierać CAŁĄ macierz — dlatego UI nie ma „pokaż wszystkie"."""
    _autorzy_nowak(30, jednostka)

    res = redaktor.get(
        reverse("bpp:zapytanie_eksport", kwargs={"export_format": "csv"}),
        {
            "model": "autor",
            "query": 'nazwisko = "Nowak"',
            "postac": "pivot",
            "pivot_row": "autor",
            "pivot_val": "liczba_autorow",
            "pivot_per_page": "25",
            "pivot_page": "2",
            "pivot_sort": "suma",
        },
    )

    wiersze = _wiersze_csv(res)
    # nagłówek + 30 autorów + RAZEM — strona 2 nie ma na to wpływu
    assert len(wiersze) == 32

"""Link do PBN w eksporcie duplikatów źródeł: root z konfiguracji uczelni.

Adres był wpisany na sztywno (``https://pbn.nauka.gov.pl/-/journal/{uid}``),
więc (a) ignorował ``pbn_api_root`` uczelni w instalacji multi-hosted,
(b) używał martwej ścieżki z czasów sedno-webapp. Kanoniczna postać to
``bpp.const.LINK_PBN_DO_ZRODLA``.
"""

import pytest
from model_bakery import baker

from bpp.models import Uczelnia
from deduplikator_zrodel.utils import _create_pbn_journal_url


@pytest.mark.django_db
def test_link_do_zrodla_uzywa_pbn_api_root_uczelni():
    uczelnia = baker.make(Uczelnia, pbn_api_root="https://pbn-a.example.com")

    url = _create_pbn_journal_url("abc123", uczelnia)

    assert url == "https://pbn-a.example.com/core/#/journal/view/abc123/current"
    assert "pbn.nauka.gov.pl" not in url
    assert "/-/journal/" not in url


@pytest.mark.django_db
def test_link_do_zrodla_bez_uczelni_jest_pusty():
    """Bez ustalonej uczelni nie zgadujemy roota — kolumna zostaje pusta."""
    baker.make(Uczelnia, pbn_api_root="https://pbn-a.example.com")
    baker.make(Uczelnia)  # >1 uczelnia → brak jednoznacznej

    assert _create_pbn_journal_url("abc123", None) == ""


def test_link_do_zrodla_bez_pbn_uid_jest_pusty():
    assert _create_pbn_journal_url(None, None) == ""


@pytest.mark.django_db
def test_eksport_xlsx_bierze_uczelnie_z_requestu(
    rf, settings, make_zrodlo, completed_scan
):
    """Pełna ścieżka: ``export_candidates_to_xlsx`` → wiersz z linkiem."""
    from openpyxl import load_workbook

    settings.ALLOWED_HOSTS = ["*"]

    from deduplikator_zrodel.utils import export_candidates_to_xlsx

    uczelnia = baker.make(Uczelnia, pbn_api_root="https://pbn-a.example.com")
    uczelnia.site.domain = "uczelnia-a.example.com"
    uczelnia.site.save()
    baker.make(Uczelnia)  # druga uczelnia → multi-hosted

    from pbn_api.models import Journal

    journal = baker.make(Journal)
    glowne = make_zrodlo(nazwa="Pismo A", pbn_uid=journal)
    duplikat = make_zrodlo(nazwa="Pismo A ", pbn_uid=None)

    scan = completed_scan()
    from deduplikator_zrodel.models import SourceDuplicateCandidate

    SourceDuplicateCandidate.objects.create(
        scan=scan,
        main_zrodlo=glowne,
        duplicate_zrodlo=duplikat,
        main_nazwa=glowne.nazwa,
        duplicate_nazwa=duplikat.nazwa,
        confidence_score=0.99,
    )

    request = rf.get("/", HTTP_HOST="uczelnia-a.example.com")
    zawartosc = export_candidates_to_xlsx(
        list(SourceDuplicateCandidate.objects.all()), request=request
    )

    from io import BytesIO

    ws = load_workbook(BytesIO(zawartosc)).active
    komorki = [c.value for row in ws.iter_rows() for c in row if c.value]
    oczekiwany = f"https://pbn-a.example.com/core/#/journal/view/{journal.pk}/current"
    assert oczekiwany in komorki
